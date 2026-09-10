"""
tests/integration/test_dieu_chuyen_hang_loat.py
===============================================
Test điều chuyển hàng loạt theo quyết định bằng Excel.

Trọng tâm là các ca mà tính năng này SINH RA để chặn — những ca đã xảy ra thật
với cách nhập lẻ (xem `docs/Fix-dieu-chuyen-don-vi/PLAN_SUA_NGAY_DIEU_CHUYEN.md`):
mã CC gõ sai, đơn vị viết tắt sai, trùng người trong một file, và quan trọng
nhất — KHÔNG được ghi một nửa rồi bỏ dở.

Chạy trên kpi_haiquan_test:
    DB_NAME=kpi_haiquan_test pytest tests/integration/test_dieu_chuyen_hang_loat.py -v

Pattern: gọi trực tiếp hàm endpoint, tự tạo/dọn dữ liệu, khôi phục hồ sơ CC sau
mỗi test.
"""
from __future__ import annotations

import io
from datetime import date, timedelta

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.api.v1.endpoints.admin_dieu_chuyen_hang_loat import (
    ghi_hang_loat,
    tai_mau_excel,
    xem_truoc,
)
from app.core.dieu_chuyen_excel import (
    SHEET_DANH_SACH,
    SHEET_HUONG_DAN,
    SHEET_MA_DON_VI,
    SHEET_NHAP_LIEU,
    TIEU_DE_NHAP_LIEU,
    TrangThaiDong,
)
from app.db.session import AsyncSessionLocal
from app.models.admin import LichSuDieuChuyen
from app.models.user_org import CongChuc, DonVi

LY_DO_TEST = "TEST dieu chuyen hang loat"


# =============================================================================
# Helpers
# =============================================================================

def _file_nhap(rows: list[tuple], sheet: str = SHEET_NHAP_LIEU) -> UploadFile:
    """Dựng file .xlsx tối thiểu giống mẫu: 1 sheet nhập liệu + tiêu đề."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(sheet)
    ws.append(TIEU_DE_NHAP_LIEU)
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return UploadFile(filename="test.xlsx", file=buf)


class _Admin:
    """Giả admin đủ dùng — endpoint chỉ đọc `.id`."""
    def __init__(self, cc_id):
        self.id = cc_id


async def _hai_cong_chuc_khac_don_vi(db):
    """Hai công chức đang hoạt động, ở HAI đơn vị khác nhau."""
    dv = (await db.execute(
        select(DonVi).where(DonVi.is_deleted == False, DonVi.ma_don_vi != "DEPT-ADMIN")
        .order_by(DonVi.ma_don_vi)
    )).scalars().all()
    assert len(dv) >= 2

    ket_qua = []
    for d in dv:
        cc = (await db.execute(
            select(CongChuc)
            .where(CongChuc.don_vi_id == d.id, CongChuc.is_active == True,
                   CongChuc.is_deleted == False, CongChuc.ma_cc != "ADMIN-001")
            .order_by(CongChuc.ma_cc).limit(1)
        )).scalar_one_or_none()
        if cc is not None:
            ket_qua.append((cc, d))
        if len(ket_qua) == 2:
            break
    assert len(ket_qua) == 2, "DB test cần 2 đơn vị mỗi bên có ít nhất 1 công chức"
    return ket_qua


async def _don_dep(db, cong_chuc_ids, don_vi_goc: dict):
    await db.execute(
        LichSuDieuChuyen.__table__.delete().where(
            LichSuDieuChuyen.cong_chuc_id.in_(cong_chuc_ids)
        )
    )
    for cc_id, dv_id in don_vi_goc.items():
        cc = (await db.execute(select(CongChuc).where(CongChuc.id == cc_id))).scalar_one()
        cc.don_vi_id = dv_id
    await db.commit()


# =============================================================================
# 1. FILE MẪU
# =============================================================================

@pytest.mark.asyncio
async def test_mau_excel_du_4_sheet_va_sheet_nhap_de_trong():
    """
    Mẫu phải có 4 sheet, và sheet nhập liệu để TRỐNG.

    Ví dụ cố ý đặt ở sheet 'Huong dan' chứ không phải sheet nhập — nếu để trên
    sheet nhập thì sẽ có lần người dùng quên xoá và nhập nhầm người mẫu vào thật.
    """
    async with AsyncSessionLocal() as db:
        resp = await tai_mau_excel(db=db, admin=_Admin(None))
        noi_dung = b"".join([chunk async for chunk in resp.body_iterator])

    wb = load_workbook(io.BytesIO(noi_dung))
    assert set(wb.sheetnames) == {
        SHEET_NHAP_LIEU, SHEET_HUONG_DAN, SHEET_DANH_SACH, SHEET_MA_DON_VI
    }

    ws = wb[SHEET_NHAP_LIEU]
    assert [c.value for c in ws[1]][:4] == TIEU_DE_NHAP_LIEU
    # Không có dòng dữ liệu nào
    co_du_lieu = [
        r for r in ws.iter_rows(min_row=2, max_col=4, values_only=True)
        if any(v not in (None, "") for v in r)
    ]
    assert co_du_lieu == [], f"Sheet nhập liệu phải trống, đang có {len(co_du_lieu)} dòng"


@pytest.mark.asyncio
async def test_mau_loc_bo_nguoi_da_nghi_va_tai_khoan_admin():
    """Sheet tra cứu chỉ liệt kê người CÒN hoạt động — để không ai điều chuyển
    nhầm người đã nghỉ."""
    async with AsyncSessionLocal() as db:
        resp = await tai_mau_excel(db=db, admin=_Admin(None))
        noi_dung = b"".join([chunk async for chunk in resp.body_iterator])

        so_active = len((await db.execute(
            select(CongChuc).where(
                CongChuc.is_active == True, CongChuc.is_deleted == False,
                CongChuc.ma_cc != "ADMIN-001",
            )
        )).scalars().all())

    ws = load_workbook(io.BytesIO(noi_dung))[SHEET_DANH_SACH]
    ma_trong_mau = [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
    assert len(ma_trong_mau) == so_active
    assert "ADMIN-001" not in ma_trong_mau


# =============================================================================
# 2. XEM TRƯỚC — nhận diện lỗi
# =============================================================================

@pytest.mark.asyncio
async def test_xem_truoc_doi_lai_ho_ten_de_phat_hien_go_nham_ma():
    """
    Đây là lưới an toàn CHÍNH của tính năng: mã CC gõ sai một chữ số vẫn có thể
    trỏ sang người thật, nên bảng xem trước phải dội lại họ tên + đơn vị hiện tại.
    """
    async with AsyncSessionLocal() as db:
        (cc1, _dv1), (_cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        f = _file_nhap([(cc1.ma_cc, dv2.ma_don_vi, date(2026, 5, 15), None)])
        kq = await xem_truoc(db=db, admin=_Admin(cc1.id), file=f)

    dong = kq["data"]["cac_dong"][0]
    assert dong["ho_ten"] == cc1.ho_ten
    assert dong["don_vi_den"] == dv2.ma_don_vi
    assert dong["trang_thai"] == TrangThaiDong.GHI_MOI
    assert kq["data"]["tom_tat"]["ghi_duoc"] is True


@pytest.mark.asyncio
async def test_xem_truoc_bat_ma_cc_khong_ton_tai():
    async with AsyncSessionLocal() as db:
        (_cc1, _dv1), (_cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        f = _file_nhap([("20ZZ-KHONG-CO", dv2.ma_don_vi, date(2026, 5, 15), None)])
        kq = await xem_truoc(db=db, admin=_Admin(None), file=f)

    dong = kq["data"]["cac_dong"][0]
    assert dong["trang_thai"] == TrangThaiDong.LOI
    assert any("không có công chức" in x for x in dong["loi"])
    assert kq["data"]["tom_tat"]["ghi_duoc"] is False


@pytest.mark.asyncio
async def test_xem_truoc_bat_don_vi_sai_va_trung_ma_trong_file():
    async with AsyncSessionLocal() as db:
        (cc1, _dv1), (_cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        f = _file_nhap([
            (cc1.ma_cc, "KHONG-CO-DON-VI-NAY", date(2026, 5, 15), None),
            (cc1.ma_cc, dv2.ma_don_vi, date(2026, 5, 15), None),
        ])
        kq = await xem_truoc(db=db, admin=_Admin(None), file=f)

    cac_dong = kq["data"]["cac_dong"]
    assert all(d["trang_thai"] == TrangThaiDong.LOI for d in cac_dong)
    assert any("đơn vị đến không hợp lệ" in x for x in cac_dong[0]["loi"])
    assert all(any("xuất hiện 2 lần" in x for x in d["loi"]) for d in cac_dong)


@pytest.mark.asyncio
async def test_xem_truoc_bo_qua_nguoi_da_o_dung_don_vi():
    """Không phải lỗi — chỉ là không có gì để làm."""
    async with AsyncSessionLocal() as db:
        (cc1, dv1), _ = await _hai_cong_chuc_khac_don_vi(db)
        f = _file_nhap([(cc1.ma_cc, dv1.ma_don_vi, date(2026, 5, 15), None)])
        kq = await xem_truoc(db=db, admin=_Admin(None), file=f)

    dong = kq["data"]["cac_dong"][0]
    assert dong["trang_thai"] == TrangThaiDong.BO_QUA
    assert kq["data"]["tom_tat"]["loi"] == 0
    # Không có dòng nào để ghi → vẫn không cho bấm ghi
    assert kq["data"]["tom_tat"]["ghi_duoc"] is False


@pytest.mark.asyncio
async def test_xem_truoc_canh_bao_nhap_muon_nhung_khong_chan():
    async with AsyncSessionLocal() as db:
        (cc1, _dv1), (_cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        ngay_cu = date.today() - timedelta(days=90)
        f = _file_nhap([(cc1.ma_cc, dv2.ma_don_vi, ngay_cu, None)])
        kq = await xem_truoc(db=db, admin=_Admin(None), file=f)

    dong = kq["data"]["cac_dong"][0]
    assert dong["trang_thai"] == TrangThaiDong.GHI_MOI      # cảnh báo, KHÔNG chặn
    assert any("nhập muộn" in x for x in dong["canh_bao"])


@pytest.mark.asyncio
async def test_xem_truoc_bao_loi_ro_khi_sai_sheet():
    """Người dùng tải nhầm file của họ thay vì file mẫu → phải nói thẳng."""
    async with AsyncSessionLocal() as db:
        f = _file_nhap([], sheet="Lam QD 2026")
        kq = await xem_truoc(db=db, admin=_Admin(None), file=f)

    loi = kq["data"]["tom_tat"]["loi_chung"]
    assert any(SHEET_NHAP_LIEU in x for x in loi)


@pytest.mark.asyncio
async def test_tu_choi_file_khong_phai_xlsx():
    async with AsyncSessionLocal() as db:
        f = UploadFile(filename="danh_sach.csv", file=io.BytesIO(b"a,b,c"))
        with pytest.raises(HTTPException) as e:
            await xem_truoc(db=db, admin=_Admin(None), file=f)
    assert e.value.status_code == 400


# =============================================================================
# 3. GHI
# =============================================================================

@pytest.mark.asyncio
async def test_ghi_ca_dot_va_sinh_lich_su_dung_ngay_quyet_dinh():
    async with AsyncSessionLocal() as db:
        (cc1, dv1), (cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        goc = {cc1.id: cc1.don_vi_id, cc2.id: cc2.don_vi_id}
        ngay = date(2026, 5, 15)

        try:
            f = _file_nhap([
                (cc1.ma_cc, dv2.ma_don_vi, ngay, "45/QĐ-HQKV8"),
                (cc2.ma_cc, dv1.ma_don_vi, ngay, "45/QĐ-HQKV8"),
            ])
            kq = await ghi_hang_loat(db=db, admin=_Admin(cc1.id), file=f)
            assert kq["data"]["so_da_ghi"] == 2

            # Đơn vị đã đổi
            for cc, dv_moi in ((cc1, dv2), (cc2, dv1)):
                moi = (await db.execute(
                    select(CongChuc).where(CongChuc.id == cc.id)
                )).scalar_one()
                assert moi.don_vi_id == dv_moi.id

            # Lịch sử ghi ĐÚNG ngày quyết định, không phải ngày hôm nay
            ls = (await db.execute(
                select(LichSuDieuChuyen).where(
                    LichSuDieuChuyen.cong_chuc_id.in_([cc1.id, cc2.id])
                )
            )).scalars().all()
            assert len(ls) == 2
            assert all(r.ngay_hieu_luc == ngay for r in ls)
            assert all(r.ngay_hieu_luc != date.today() for r in ls)
            assert all("45/QĐ-HQKV8" in (r.ly_do or "") for r in ls)
        finally:
            await _don_dep(db, [cc1.id, cc2.id], goc)


@pytest.mark.asyncio
async def test_mot_dong_loi_thi_tu_choi_ca_dot_khong_ghi_mot_nua():
    """
    Ca quan trọng nhất. Ghi một nửa rồi báo "59/62 thành công" là đẩy việc dò
    tìm 3 dòng còn lại sang người dùng — đúng kiểu lỗi mà tính năng này sinh ra
    để chặn.
    """
    async with AsyncSessionLocal() as db:
        (cc1, _dv1), (_cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        # Giữ giá trị NGUYÊN THỦY: sau `rollback` mọi đối tượng ORM bị hết hạn,
        # chạm vào thuộc tính sẽ kích hoạt lazy-load ngoài greenlet → MissingGreenlet.
        cc1_id, cc1_ma, don_vi_truoc = cc1.id, cc1.ma_cc, cc1.don_vi_id
        dv2_ma = dv2.ma_don_vi

        try:
            f = _file_nhap([
                (cc1_ma, dv2_ma, date(2026, 5, 15), None),               # hợp lệ
                ("20ZZ-KHONG-CO", dv2_ma, date(2026, 5, 15), None),      # hỏng
            ])
            with pytest.raises(HTTPException) as e:
                await ghi_hang_loat(db=db, admin=_Admin(cc1_id), file=f)
            assert e.value.status_code == 400

            # Dòng hợp lệ cũng KHÔNG được ghi
            await db.rollback()
            moi = (await db.execute(
                select(CongChuc).where(CongChuc.id == cc1_id)
            )).scalar_one()
            assert moi.don_vi_id == don_vi_truoc

            ls = (await db.execute(
                select(LichSuDieuChuyen).where(LichSuDieuChuyen.cong_chuc_id == cc1_id)
            )).scalars().all()
            assert ls == []
        finally:
            await _don_dep(db, [cc1_id], {cc1_id: don_vi_truoc})


@pytest.mark.asyncio
async def test_ghi_nhieu_dot_khac_ngay_trong_mot_file():
    """File của TCCB chứa nhiều đợt cùng lúc — mỗi dòng một ngày riêng."""
    async with AsyncSessionLocal() as db:
        (cc1, dv1), (cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        goc = {cc1.id: cc1.don_vi_id, cc2.id: cc2.don_vi_id}

        try:
            f = _file_nhap([
                (cc1.ma_cc, dv2.ma_don_vi, date(2026, 2, 4), None),
                (cc2.ma_cc, dv1.ma_don_vi, date(2026, 7, 3), None),
            ])
            await ghi_hang_loat(db=db, admin=_Admin(cc1.id), file=f)

            ngay_theo_cc = {
                r.cong_chuc_id: r.ngay_hieu_luc
                for r in (await db.execute(
                    select(LichSuDieuChuyen).where(
                        LichSuDieuChuyen.cong_chuc_id.in_([cc1.id, cc2.id])
                    )
                )).scalars().all()
            }
            assert ngay_theo_cc[cc1.id] == date(2026, 2, 4)
            assert ngay_theo_cc[cc2.id] == date(2026, 7, 3)
        finally:
            await _don_dep(db, [cc1.id, cc2.id], goc)


@pytest.mark.asyncio
async def test_ngay_dang_chuoi_ddmmyyyy_van_doc_duoc():
    """Người dùng gõ tay vào ô chưa định dạng ngày thì Excel lưu thành chuỗi."""
    async with AsyncSessionLocal() as db:
        (cc1, _dv1), (_cc2, dv2) = await _hai_cong_chuc_khac_don_vi(db)
        f = _file_nhap([(cc1.ma_cc, dv2.ma_don_vi, "15/05/2026", None)])
        kq = await xem_truoc(db=db, admin=_Admin(None), file=f)

    dong = kq["data"]["cac_dong"][0]
    assert dong["trang_thai"] == TrangThaiDong.GHI_MOI
    assert dong["ngay_hieu_luc"] == "2026-05-15"
